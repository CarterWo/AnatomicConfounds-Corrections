"""Builds Supplementary Table 3 (variance components) from the per-gene result CSVs.

usage: python3 make_supp_table.py <results_dir> <out.xlsx>

The first sheet summarizes each component across genes (median, IQR, 90th percentile,
mean) as a percentage of total variance; the other sheets list the per-gene values.
"""
import csv
import statistics as st
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

VC_DIR, OUT = sys.argv[1], sys.argv[2]

SETS = [
    ("CosMx (this study)", "Napari (4 broad regions + unassigned)", "vc_cosmx_napari.csv", "Per-gene CosMx Napari"),
    ("CosMx (this study)", "QUINT (36 atlas regions)", "vc_cosmx_quint.csv", "Per-gene CosMx QUINT"),
    ("MERFISH (Zhuang-ABCA-1-4)", "Napari-equivalent (CCFv3 division)", "vc_merfish_napari.csv", "Per-gene MERFISH Napari-eq"),
    ("MERFISH (Zhuang-ABCA-1-4)", "QUINT-equivalent (CCFv3 structure)", "vc_merfish_quint.csv", "Per-gene MERFISH QUINT-eq"),
]
COMPS = [("cell_type", "Cell type"), ("region", "Anatomic region"), ("sample_ID", "Biological sample"),
         ("Treatment", "Treatment group"), ("log_depth", "Log unique genes count"), ("Residuals", "Residual")]


def percentile(values, p):
    # linear interpolation, same as R's default quantile()
    values = sorted(values)
    pos = (len(values) - 1) * p
    lo = int(pos)
    hi = min(lo + 1, len(values) - 1)
    return values[lo] + (values[hi] - values[lo]) * (pos - lo)


thin = Side(style="thin")
hdr_fill = PatternFill("solid", fgColor="FFD9D9D9")
band_fill = PatternFill("solid", fgColor="FFF7F7F7")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Table S3"
hdr = ["Dataset", "Region labels", "Genes", "Variance component", "Median (%)",
       "25th percentile (%)", "75th percentile (%)", "90th percentile (%)", "Mean (%)"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(1, j, h)
    c.font = Font(name="Calibri", size=11, bold=True)
    c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = Border(bottom=thin)

row, band, per = 2, False, {}
for ds, lab, fn, sheet in SETS:
    rows = list(csv.DictReader(open(f"{VC_DIR}/{fn}")))
    per[sheet] = rows
    band = not band
    for key, name in COMPS:
        if key not in rows[0]:
            continue
        v = [100 * float(r[key]) for r in rows]
        vals = [ds, lab, len(rows), name, round(percentile(v, .5), 2), round(percentile(v, .25), 2),
                round(percentile(v, .75), 2), round(percentile(v, .9), 2), round(st.mean(v), 2)]
        for j, val in enumerate(vals, 1):
            c = ws.cell(row, j, val)
            c.font = Font(name="Calibri", size=11)
            if band:
                c.fill = band_fill
        row += 1
for col, w in zip("ABCDEFGHI", [28, 36, 8, 24, 12, 13, 13, 13, 10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"

row += 1
cap = (
    "Supplementary Table 3. Estimated variance components for model covariates, summarized "
    "across genes in the first sheet, with separate per-gene sheets for the CosMx dataset with "
    "Napari and QUINT labels and for the MERFISH dataset with Napari- and QUINT-equivalent "
    "labels."
)
c = ws.cell(row, 1, cap)
c.font = Font(name="Calibri", size=9)
c.alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
ws.row_dimensions[row].height = 45

for sheet, rows in per.items():
    s = wb.create_sheet(sheet)
    keys = [k for k, _ in COMPS if k in rows[0]]
    heads = ["Gene"] + [n + " (%)" for k, n in COMPS if k in rows[0]]
    for j, h in enumerate(heads, 1):
        c = s.cell(1, j, h)
        c.font = Font(bold=True)
        c.fill = hdr_fill
        c.border = Border(bottom=thin)
    for i, r in enumerate(sorted(rows, key=lambda r: -float(r["region"])), 2):
        s.cell(i, 1, r["gene"])
        for j, k in enumerate(keys, 2):
            s.cell(i, j, round(100 * float(r[k]), 3))
    s.freeze_panes = "B2"
    s.column_dimensions["A"].width = 16
wb.save(OUT)
print("saved", OUT)
